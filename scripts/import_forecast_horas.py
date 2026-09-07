"""
import_forecast_combinado.py
----------------------------
Importa el forecast desde las DOS fuentes que definio el negocio:

    Forecast S&P (9-1), hoja 'Forecast Update'
        -> S&P legacy: SO / PR / Tools, en ARG, MX y CR
        -> el offering sale de la columna 'Offering Label'
        -> HL y SL se distinguen por el color de la celda CHG

    01 - FORECAST CONTROL S&P.xlsm, hoja 'S&P'
        -> Digi: S4 / Ariba / Oracle
        -> el offering sale de la columna 'Sub-Offerings'
        -> HL y SL vienen en columnas explicitas, no por color

El 9-1 no trae a la gente de Digi y el CONTROL clasifica distinto a los legacy
(los 17 que el 9-1 pone en Tools el CONTROL los pone en PR). Por eso cada grupo
se toma de su fuente y no se mezclan: para un mismo EID manda una sola planilla.

Uso:
    python import_forecast_combinado.py --dry-run
    python import_forecast_combinado.py
    python import_forecast_combinado.py --update-location
"""

import argparse
import asyncio
import logging
import os
import re
import sys
from collections import Counter, defaultdict

import asyncpg
from dotenv import load_dotenv
from openpyxl import load_workbook

# El .env vive en la raiz del repo, no en scripts/
load_dotenv(os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '.env'))

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s',
    datefmt='%H:%M:%S',
)
log = logging.getLogger(__name__)

AZURE = dict(
    host=os.getenv('DB_HOST'),
    port=int(os.getenv('DB_PORT', 5432)),
    user=os.getenv('DB_USER'),
    password=os.getenv('DB_PASSWORD'),
    database=os.getenv('DB_NAME'),
    ssl='require',
)

DESKTOP = os.path.join(os.path.expanduser('~'), 'OneDrive - Accenture', 'Desktop')
DEFAULT_LEGACY = os.path.join(DESKTOP, '2026 9 1 - Forecast S&P.xlsx')
DEFAULT_DIGI = os.path.join(DESKTOP, '01 - FORECAST CONTROL S&P (1).xlsm')

# ---------------------------------------------------------------- 9-1 (legacy)

SH_LEGACY = 'Forecast Update'
LG_HEADER_ROW = 2        # fila con los labels de periodo
LG_SUBHEADER_ROW = 10    # fila con CHG / SAH / CHG (%)
LG_FIRST_DATA_ROW = 11
LG_COL_LOCATION = 1
LG_COL_EID = 2
LG_COL_OFFERING = 3
LG_COL_STATUS = 6

# El bloque vigente se corre a la derecha cada mes. Los labels se repiten entre
# el bloque viejo y el nuevo, asi que hay que decirle desde donde leer. Se
# expone como flag para no tener que editar el codigo el mes que viene.
LG_MIN_PERIOD_COL = 82

OFFERINGS_LEGACY = {'SO', 'PR', 'Tools'}
LOCATIONS_LEGACY = {'ARG', 'MX', 'CR'}

# Colores del 9-1. La leyenda vive en 'Forecast Update' B93:B96.
COLOR_SL = {
    'FFFFEB9C',      # Assumptions ISG Assessment
    'FFFFCCCC',      # Assumptions No R
    'FFFFC7CE',      # variante rosa de No R
    'THEME4+0.8',    # Assumptions R
}
COLOR_SUBKIND = {
    'FFFFEB9C': 'isg_assessment',
    'FFFFCCCC': 'no_r',
    'FFFFC7CE': 'no_r',
    'THEME4+0.8': 'r',
}
COLOR_PPA = 'FFCC66FF'   # cascadeo de horas: el CHG de la celda sigue siendo HL

PERIOD_LABEL_MAP = {
    'ENE P1': 'Ene-P1', 'ENE P2': 'Ene-P2', 'FEB P1': 'Feb-P1', 'FEB P2': 'Feb-P2',
    'MAR P1': 'Mar-P1', 'MAR P2': 'Mar-P2', 'ABR P1': 'Abr-P1', 'ABR P2': 'Abr-P2',
    'MAY P1': 'May-P1', 'MAY P2': 'May-P2', 'JUN P1': 'Jun-P1', 'JUN P2': 'Jun-P2',
    'JUL P1': 'Jul-P1', 'JUL P2': 'Jul-P2', 'AGO P1': 'Ago-P1', 'AGO P2': 'Ago-P2',
    'SEP P1': 'Sep-P1', 'SEP P2': 'Sep-P2', 'OCT P1': 'Oct-P1', 'OCT P2': 'Oct-P2',
    'NOV P1': 'Nov-P1', 'NOV P2': 'Nov-P2', 'DIC P1': 'Dic-P1', 'DIC P2': 'Dic-P2',
}

# ----------------------------------------------------------- CONTROL (digi)

SH_DIGI = 'S&P'
DG_HEADER_ROW = 1
DG_FIRST_DATA_ROW = 2
DG_COL_STATUS = 1
DG_COL_SUBOFFERING = 3
DG_COL_EID = 12

OFFERINGS_DIGI = {'S4', 'Ariba', 'Oracle'}

MES_EN_ES = {
    'JAN': 'Ene', 'FEB': 'Feb', 'MAR': 'Mar', 'APR': 'Abr', 'MAY': 'May',
    'JUN': 'Jun', 'JUL': 'Jul', 'AUG': 'Ago', 'SEP': 'Sep', 'OCT': 'Oct',
    'NOV': 'Nov', 'DEC': 'Dic',
}

RE_DG_HL = re.compile(r'^CHG\s+HL\s+([A-Z]{3})\s+P([12])\s+(\d{4})$', re.I)
RE_DG_SL = re.compile(r'^CHG\s+SL\s+([A-Z]{3})\s+P([12])\s+(\d{4})$', re.I)
RE_DG_SAH = re.compile(r'^SAH\s+P[12]\s+([A-Z]{3})\s+P([12])\s+(\d{4})$', re.I)
RE_DG_PCT = re.compile(r'^CHG%\s+([A-Z]{3})\s+P([12])\s+(\d{4})$', re.I)

LOCATION_A_COUNTRY = {'ARG': 'Argentina', 'MX': 'Mexico', 'CR': 'Costa Rica'}


def num(v):
    """None si la celda no tiene un numero usable. No confunde 0 con vacio."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def cell_color(cell) -> str:
    """Color de relleno. Soporta rgb, theme e indexed: los 'Assumptions R' usan
    theme4+0.8 y leyendo solo rgb se veian como celda sin color, o sea HL."""
    f = cell.fill
    if not f or not f.patternType:
        return 'NOFILL'
    fg = f.fgColor
    if not fg:
        return 'NOFILL'
    try:
        if fg.type == 'rgb' and fg.rgb and str(fg.rgb).upper() != '00000000':
            return str(fg.rgb).upper()
        if fg.type == 'theme':
            return f'THEME{fg.theme}+{round(float(fg.tint), 3)}'
        if fg.type == 'indexed':
            return f'INDEXED{fg.indexed}'
    except Exception:
        pass
    return 'NOFILL'


# =============================================================== 9-1 legacy

def build_triplets(ws):
    """Tripletes (chg, sah, pct) leidos de la fila de subcabeceras.

    No se puede asumir que la columna del label del periodo sea la del CHG: hay
    periodos donde el label esta corrido y cae sobre el SAH, lo que daba
    porcentajes de 8000%.
    """
    tipo = {}
    for cell in ws[LG_SUBHEADER_ROW]:
        if not isinstance(cell.value, str):
            continue
        h = cell.value.strip().upper()
        if '%' in h:
            tipo[cell.column] = 'PCT'
        elif h.startswith('CHG'):
            tipo[cell.column] = 'CHG'
        elif h.startswith('SAH'):
            tipo[cell.column] = 'SAH'

    triplets = []
    for col in sorted(tipo):
        if tipo[col] != 'CHG':
            continue
        sah = next((c for c in sorted(tipo) if c > col and tipo[c] == 'SAH'), None)
        pct = next((c for c in sorted(tipo) if c > col and tipo[c] == 'PCT'), None)
        if sah and pct and sah < pct:
            triplets.append((col, sah, pct))
    return triplets


def parse_legacy(path: str, min_period_col: int):
    log.info(f'[9-1] Leyendo {path}')
    wb = load_workbook(path, read_only=False, data_only=True)
    if SH_LEGACY not in wb.sheetnames:
        raise SystemExit(f"[9-1] No existe la hoja '{SH_LEGACY}'. Hojas: {wb.sheetnames}")
    ws = wb[SH_LEGACY]

    if (ws.cell(LG_SUBHEADER_ROW, LG_COL_EID).value or '').strip() != 'EID':
        raise SystemExit(
            f'[9-1] La col {LG_COL_EID} de la fila {LG_SUBHEADER_ROW} deberia ser '
            f'"EID" y es "{ws.cell(LG_SUBHEADER_ROW, LG_COL_EID).value}". '
            'El layout cambio, revisar las constantes LG_COL_*.')

    # Los labels se repiten entre bloques; gana la ultima aparicion a la
    # derecha del corte.
    period_cols, descartados = {}, defaultdict(list)
    for cell in ws[LG_HEADER_ROW]:
        if not isinstance(cell.value, str):
            continue
        key = cell.value.strip().upper()
        if key not in PERIOD_LABEL_MAP:
            continue
        periodo = PERIOD_LABEL_MAP[key]
        if cell.column < min_period_col:
            descartados[periodo].append(cell.column)
            continue
        period_cols[periodo] = cell.column

    if not period_cols:
        raise SystemExit(
            f'[9-1] No se detecto ningun periodo a partir de la columna '
            f'{min_period_col}. Revisar --min-period-col.')

    solo_viejos = sorted(p for p in descartados if p not in period_cols)
    if solo_viejos:
        log.info(f'[9-1] Periodos que solo existen en el bloque anterior al corte, '
                 f'se ignoran: {solo_viejos}')

    triplets = build_triplets(ws)
    period_tri = {}
    for periodo, lcol in period_cols.items():
        t = next((x for x in triplets if x[0] <= lcol <= x[2]), None)
        if t is None:
            log.warning(f'[9-1] {periodo}: sin triplete CHG/SAH/CHG%, se omite')
            continue
        period_tri[periodo] = t
        if t[0] != lcol:
            log.warning(f'[9-1] {periodo}: el label esta en col {lcol} pero el CHG '
                        f'es col {t[0]} (corregido)')

    log.info(f'[9-1] Periodos vigentes: {len(period_tri)} -> {sorted(period_tri)}')

    empleados, datos = {}, defaultdict(dict)
    ppa_cells, descuadres, sin_color = [], [], Counter()
    fuera_de_alcance = Counter()

    for r in range(LG_FIRST_DATA_ROW, ws.max_row + 1):
        eid = ws.cell(r, LG_COL_EID).value
        if not isinstance(eid, str) or not eid.strip() or eid.startswith('='):
            continue
        eid = eid.strip().lower()

        loc = ws.cell(r, LG_COL_LOCATION).value
        loc = loc.strip().upper() if isinstance(loc, str) else ''
        off = ws.cell(r, LG_COL_OFFERING).value
        off = off.strip() if isinstance(off, str) else ''
        reserva = ws.cell(r, LG_COL_STATUS).value
        reserva = reserva.strip() if isinstance(reserva, str) else None

        # El bloque de totales del Excel cuenta Location ARG/MX/CR con offering
        # SO/PR/Tools. El resto son filas de leyenda, placeholders 'DIGI' y
        # sobrantes: quedan fuera a proposito.
        if loc not in LOCATIONS_LEGACY or off not in OFFERINGS_LEGACY:
            fuera_de_alcance[f'{loc or "(sin loc)"} / {off or "(sin offering)"}'] += 1
            continue

        empleados[eid] = {'location': loc, 'offering': off, 'reserva_status': reserva, 'fuente': '9-1'}

        for periodo, (ccol, scol, pcol) in period_tri.items():
            chg_cell = ws.cell(r, ccol)
            chg = num(chg_cell.value)
            sah = num(ws.cell(r, scol).value)
            if sah is None:
                continue

            color = cell_color(chg_cell)
            if color == COLOR_PPA:
                ppa_cells.append((eid, periodo, chg_cell.coordinate, chg))

            es_sl = color in COLOR_SL
            hl = 0.0 if es_sl else (chg or 0.0)
            sl = (chg or 0.0) if es_sl else 0.0

            if chg and color not in COLOR_SL and color not in (
                    'NOFILL', 'THEME0+0.0', 'THEME0+-0.15', COLOR_PPA):
                sin_color[color] += 1

            # Chequeo contra el CHG% del propio Excel. Si no coincide estamos
            # leyendo columnas de periodos distintos: mejor verlo aca que
            # descubrirlo despues en la app.
            pct_excel = num(ws.cell(r, pcol).value)
            if pct_excel is not None and sah > 0:
                pct_calc = (hl + sl) / sah
                if abs(pct_excel - pct_calc) > 0.011:
                    descuadres.append((eid, periodo, hl, sl, sah,
                                       round(pct_calc * 100, 1), round(pct_excel * 100, 1)))

            datos[eid][periodo] = {
                'hl': hl, 'sl': sl, 'sah': sah,
                'subkind': COLOR_SUBKIND.get(color) if es_sl else None,
            }

    wb.close()

    if fuera_de_alcance:
        log.info(f'[9-1] Filas fuera de alcance (no son ARG/MX/CR + SO/PR/Tools): '
                 f'{dict(fuera_de_alcance)}')
    if sin_color:
        log.warning(f'[9-1] Colores no clasificados, se tomaron como HL: {dict(sin_color)}')
    if ppa_cells:
        log.info(f'[9-1] {len(ppa_cells)} celdas de cascadeo (lila): el CHG se carga '
                 'como HL. El detalle de PPA lo maneja import_from_excel.')
    if descuadres:
        log.warning(f'[9-1] {len(descuadres)} celdas donde el CHG% calculado no '
                    'coincide con el del Excel:')
        for d in descuadres[:8]:
            log.warning(f'    {d[0]} {d[1]} hl={d[2]} sl={d[3]} sah={d[4]} '
                        f'-> calc {d[5]}% vs excel {d[6]}%')

    log.info(f'[9-1] {len(empleados)} empleados legacy · '
             f'{Counter((v["location"], v["offering"]) for v in empleados.values())}')
    return empleados, datos, sorted(period_tri)


# ============================================================ CONTROL digi

def parse_digi(path: str):
    log.info(f'[CONTROL] Leyendo {path}')
    wb = load_workbook(path, read_only=False, data_only=True)
    if SH_DIGI not in wb.sheetnames:
        raise SystemExit(f"[CONTROL] No existe la hoja '{SH_DIGI}'. Hojas: {wb.sheetnames}")
    ws = wb[SH_DIGI]

    if (ws.cell(DG_HEADER_ROW, DG_COL_EID).value or '').strip() != 'Enterprise ID':
        raise SystemExit(
            f'[CONTROL] La col {DG_COL_EID} deberia ser "Enterprise ID" y es '
            f'"{ws.cell(DG_HEADER_ROW, DG_COL_EID).value}". El layout cambio.')

    # Las columnas se detectan por encabezado, asi que el bloque puede crecer
    # sin tocar el codigo.
    cols = defaultdict(dict)
    for cell in ws[DG_HEADER_ROW]:
        if not isinstance(cell.value, str):
            continue
        h = cell.value.strip()
        for rx, key in ((RE_DG_HL, 'hl'), (RE_DG_SL, 'sl'),
                        (RE_DG_SAH, 'sah'), (RE_DG_PCT, 'pct')):
            m = rx.match(h)
            if not m:
                continue
            mes = MES_EN_ES.get(m.group(1).upper())
            if mes:
                cols[f'{mes}-P{m.group(2)}'][key] = cell.column
            break

    period_cols = {p: c for p, c in cols.items()
                   if all(k in c for k in ('hl', 'sl', 'sah'))}
    incompletos = {p: sorted(c) for p, c in cols.items() if p not in period_cols}
    if incompletos:
        log.warning(f'[CONTROL] Periodos con columnas incompletas, se omiten: {incompletos}')
    if not period_cols:
        raise SystemExit('[CONTROL] No se detecto ningun periodo en los encabezados.')

    log.info(f'[CONTROL] Periodos: {len(period_cols)} -> {sorted(period_cols)}')

    empleados, datos = {}, defaultdict(dict)
    descuadres, inactivos, no_digi = [], 0, Counter()

    for r in range(DG_FIRST_DATA_ROW, ws.max_row + 1):
        eid = ws.cell(r, DG_COL_EID).value
        if not isinstance(eid, str) or not eid.strip() or eid.startswith('='):
            continue
        eid = eid.strip().lower()

        estado = ws.cell(r, DG_COL_STATUS).value
        if isinstance(estado, str) and estado.strip() and estado.strip().lower() != 'active':
            inactivos += 1
            continue

        sub = ws.cell(r, DG_COL_SUBOFFERING).value
        sub = sub.strip() if isinstance(sub, str) else ''
        if sub not in OFFERINGS_DIGI:
            no_digi[sub or '(vacio)'] += 1
            continue

        empleados[eid] = {'location': 'ARG', 'offering': sub, 'reserva_status': None, 'fuente': 'CONTROL'}

        for periodo, c in period_cols.items():
            hl = num(ws.cell(r, c['hl']).value) or 0.0
            sl = num(ws.cell(r, c['sl']).value) or 0.0
            sah = num(ws.cell(r, c['sah']).value)
            if sah is None:
                continue

            pct_excel = num(ws.cell(r, c['pct']).value) if 'pct' in c else None
            if pct_excel is not None and sah > 0:
                pct_calc = (hl + sl) / sah
                if abs(pct_excel - pct_calc) > 0.011:
                    descuadres.append((eid, periodo, hl, sl, sah,
                                       round(pct_calc * 100, 1), round(pct_excel * 100, 1)))

            datos[eid][periodo] = {'hl': hl, 'sl': sl, 'sah': sah, 'subkind': None}

    wb.close()

    if inactivos:
        log.info(f'[CONTROL] {inactivos} filas descartadas por status distinto de Active')
    if no_digi:
        log.info(f'[CONTROL] Sub-Offerings que no son Digi, se ignoran '
                 f'(vienen del 9-1): {dict(no_digi)}')
    if descuadres:
        log.warning(f'[CONTROL] {len(descuadres)} celdas con CHG% que no coincide:')
        for d in descuadres[:8]:
            log.warning(f'    {d[0]} {d[1]} hl={d[2]} sl={d[3]} sah={d[4]} '
                        f'-> calc {d[5]}% vs excel {d[6]}%')

    log.info(f'[CONTROL] {len(empleados)} empleados Digi · '
             f'{Counter(v["offering"] for v in empleados.values())}')
    return empleados, datos, sorted(period_cols)


# ==================================================================== reporte

ORDEN_OFFERING = ['SO', 'PR', 'Tools', 'S4', 'Ariba', 'Oracle']


def reporte(empleados, datos, periodo):
    """Totales por pais y offering, en el mismo corte que el bloque del Excel."""
    agg = defaultdict(lambda: {'hl': 0.0, 'sl': 0.0, 'sah': 0.0, 'hc': 0})
    for eid, info in empleados.items():
        d = datos[eid].get(periodo)
        k = (info['location'], info['offering'])
        a = agg[k]
        a['hc'] += 1
        if d:
            a['hl'] += d['hl']
            a['sl'] += d['sl']
            a['sah'] += d['sah']

    def linea(label, a, indent=''):
        net = a['hl'] + a['sl']
        pct = net / a['sah'] * 100 if a['sah'] else 0
        print(f"  {indent}{label:22} {a['hc']:4} {a['hl']:8.0f} {a['sl']:7.0f} "
              f"{net:8.0f} {a['sah']:7.0f} {pct:7.1f}%")

    print()
    print(f'  === {periodo} · totales (comparar contra el bloque del Excel) ===')
    print(f"  {'grupo':22} {'HC':>4} {'HL':>8} {'SL':>7} {'NET':>8} {'SAH':>7} {'CHG%':>8}")
    print('  ' + '-' * 68)

    total = {'hl': 0.0, 'sl': 0.0, 'sah': 0.0, 'hc': 0}
    for loc, nombre in (('ARG', 'Total S&P Arg'), ('MX', 'Total S&P Mexico'),
                        ('CR', 'Total S&P Costa Rica')):
        offs = [(o, a) for (l, o), a in agg.items() if l == loc]
        if not offs:
            continue
        sub = {'hl': 0.0, 'sl': 0.0, 'sah': 0.0, 'hc': 0}
        for _o, a in offs:
            for k in sub:
                sub[k] += a[k]
        linea(nombre, sub)
        for o in ORDEN_OFFERING:
            for oo, a in offs:
                if oo == o:
                    linea(o, a, indent='  ')
        for k in total:
            total[k] += sub[k]

    print('  ' + '-' * 68)
    linea('TOTAL', total)
    print()


# ==================================================================== escritura

async def write_db(empleados, datos, periodos_por_fuente, dry_run, update_location):
    conn = await asyncpg.connect(**AZURE)
    try:
        db_periods = {r['period_name']: (r['start_date'], r['end_date'])
                      for r in await conn.fetch(
                          'SELECT period_name, start_date, end_date FROM periods')}
        db_emp = {r['eid'].strip().lower(): r for r in await conn.fetch(
            'SELECT eid, name, offering, location, country, active, reserva_status FROM employees')}
        log.info(f'DB: {len(db_periods)} periodos · {len(db_emp)} empleados')

        faltan_en_db = sorted(set(empleados) - set(db_emp))
        sobran_en_db = sorted(
            e for e, r in db_emp.items() if r['active'] and e not in empleados)

        if faltan_en_db:
            log.warning(f'{len(faltan_en_db)} EIDs de los Excel que no existen en '
                        f'employees, se omiten: {faltan_en_db}')
        if sobran_en_db:
            log.warning(f'{len(sobran_en_db)} empleados activos en la DB que no estan '
                        f'en ninguna de las dos fuentes: {sobran_en_db}')

        # --- offering, location y reserva_status
        cambios_emp, cambios_fu, cambios_loc, cambios_reserva = [], [], [], []
        fu_actual = {r['eid'].strip().lower(): r['offering'] for r in await conn.fetch(
            """
            SELECT DISTINCT ON (eid) eid, offering FROM forecast_update
            ORDER BY eid, updated_at DESC NULLS LAST
            """)}

        for eid, info in empleados.items():
            row = db_emp.get(eid)
            if row is None:
                continue
            if (row['offering'] or '') != info['offering']:
                cambios_emp.append((eid, row['name'], row['offering'], info['offering']))
            if (fu_actual.get(eid) or '') != info['offering']:
                cambios_fu.append((eid, row['name'], fu_actual.get(eid), info['offering']))
            pais = LOCATION_A_COUNTRY.get(info['location'])
            if (row['location'] or '') != info['location'] or (row['country'] or '') != pais:
                cambios_loc.append((eid, row['location'], row['country'],
                                    info['location'], pais))
            nueva_reserva = info.get('reserva_status')
            if (row['reserva_status'] or None) != (nueva_reserva or None):
                cambios_reserva.append((eid, row['name'], row['reserva_status'], nueva_reserva))

        def tabla(titulo, filas, cols=4):
            if not filas:
                return
            print(f'  --- {titulo} ({len(filas)}) ---')
            for f in filas[:40]:
                if cols == 4:
                    print(f'    {f[0]:26} {str(f[2] or "-"):>10}  ->  {f[3]}')
                else:
                    print(f'    {f[0]:26} {str(f[1]):>6}/{str(f[2]):<12}  ->  '
                          f'{f[3]}/{f[4]}')
            if len(filas) > 40:
                print(f'    ... y {len(filas) - 40} mas')
            print()

        print()
        print('  === cambios de clasificacion ===')
        tabla('employees.offering', cambios_emp)
        tabla('forecast_update.offering', cambios_fu)
        tabla('location / country', cambios_loc, cols=5)
        tabla('reserva_status', cambios_reserva)
        if not update_location and cambios_loc:
            print('    (location/country NO se tocan sin --update-location)')
            print()

        # --- bloques de chargeability
        blocks, sah_rows = [], []
        periodos_tocados = set()
        for eid, info in empleados.items():
            if eid not in db_emp:
                continue
            for periodo, d in datos[eid].items():
                if periodo not in db_periods:
                    continue
                periodos_tocados.add((eid, periodo))
                sah_rows.append((eid, periodo, float(d['sah'])))
                if d['sah'] <= 0:
                    continue
                start, end = db_periods[periodo]
                for horas, scenario in ((d['hl'], 'effective'), (d['sl'], 'assumption')):
                    if horas <= 0:
                        continue
                    blocks.append((
                        eid, periodo, round(horas / d['sah'] * 100, 2), scenario,
                        d['subkind'] if scenario == 'assumption' else None, start, end))

        hl_n = sum(1 for b in blocks if b[3] == 'effective')
        log.info(f'Bloques a escribir: {len(blocks)} (HL {hl_n} / SL {len(blocks) - hl_n})')
        log.info(f'SAH a escribir: {len(sah_rows)} pares (eid, periodo)')

        if dry_run:
            print('  DRY RUN: no se escribio nada en la DB.')
            print()
            return

        claves = sorted(periodos_tocados)
        periodos_db = sorted({p for _e, p in claves})

        async with conn.transaction():
            if cambios_emp:
                await conn.executemany(
                    'UPDATE employees SET offering = $2 WHERE eid = $1',
                    [(e, n) for e, _nm, _a, n in cambios_emp])

            if cambios_reserva:
                await conn.executemany(
                    'UPDATE employees SET reserva_status = $2 WHERE eid = $1',
                    [(e, n) for e, _nm, _a, n in cambios_reserva])

            if cambios_fu:
                # forecast_update puede tener varias filas por eid. Se actualizan
                # todas para que el DISTINCT ON de latest_fu devuelva el valor
                # nuevo sin importar cual quede primera.
                await conn.execute(
                    """
                    WITH nuevos(eid, offering) AS (
                        SELECT * FROM UNNEST($1::text[], $2::text[])
                    )
                    UPDATE forecast_update fu
                    SET offering = n.offering, updated_at = NOW()
                    FROM nuevos n
                    WHERE fu.eid = n.eid AND COALESCE(fu.offering, '') <> n.offering
                    """,
                    [e for e, _nm, _a, _n in cambios_fu],
                    [n for _e, _nm, _a, n in cambios_fu])

                # Los que no tienen ninguna fila en forecast_update no entran al
                # GROUP BY de los totales y su SAH queda fuera del denominador.
                sin_fila = [(e, n) for e, _nm, _a, n in cambios_fu if e not in fu_actual]
                if sin_fila:
                    await conn.executemany(
                        'INSERT INTO forecast_update (eid, offering, updated_at) '
                        'VALUES ($1, $2, NOW())', sin_fila)
                    log.info(f'{len(sin_fila)} filas nuevas en forecast_update')

            if update_location and cambios_loc:
                await conn.executemany(
                    'UPDATE employees SET location = $2, country = $3 WHERE eid = $1',
                    [(e, l, p) for e, _al, _ap, l, p in cambios_loc])

            # El DELETE va por (eid, periodo) sin filtrar scenario_type: si una
            # celda pasa de HL a SL entre corridas, filtrar por tipo dejaria viva
            # la fila anterior y el CHG neto se duplicaria.
            log.info(f'Limpiando {len(claves)} combinaciones (eid, periodo)...')
            await conn.executemany(
                'DELETE FROM chargeability_blocks WHERE eid = $1 AND period_name = $2',
                claves)

            log.info(f'Insertando {len(blocks)} bloques...')
            await conn.executemany(
                'INSERT INTO chargeability_blocks '
                '(eid, period_name, chargeability_pct, scenario_type, '
                ' assumption_kind, start_date, end_date, created_at) '
                'VALUES ($1, $2, $3, $4, $5, $6, $7, NOW())',
                blocks)

            log.info(f'Actualizando SAH de {len(sah_rows)} filas...')
            await conn.executemany(
                'INSERT INTO forecast_periods (eid, period_name, sah) '
                'VALUES ($1, $2, $3) '
                'ON CONFLICT (eid, period_name) DO UPDATE SET sah = EXCLUDED.sah',
                sah_rows)

            # Rederiva el CHG con el SAH nuevo. Misma formula que
            # recalculate_service.recalculate_employee, en un solo statement.
            await conn.execute(
                """
                WITH totals AS (
                    SELECT eid, period_name,
                           COALESCE(SUM(chargeability_pct)
                               FILTER (WHERE scenario_type = 'effective'),  0) AS hl_pct,
                           COALESCE(SUM(chargeability_pct)
                               FILTER (WHERE scenario_type = 'assumption'), 0) AS sl_pct
                    FROM chargeability_blocks
                    WHERE period_name = ANY($1)
                    GROUP BY eid, period_name
                )
                UPDATE forecast_periods fp
                SET chg_pct_hl = t.hl_pct,
                    chg_pct_sl = t.sl_pct,
                    chg_hl     = ROUND(fp.sah * t.hl_pct / 100.0),
                    chg_sl     = ROUND(fp.sah * t.sl_pct / 100.0),
                    chg        = ROUND(fp.sah * (t.hl_pct + t.sl_pct) / 100.0)
                FROM totals t
                WHERE fp.eid = t.eid AND fp.period_name = t.period_name
                """,
                periodos_db)

            # Sin bloques el empleado no tiene horas cargables en ese periodo.
            # Si no se hace explicito queda el CHG de la corrida anterior.
            await conn.execute(
                """
                UPDATE forecast_periods fp
                SET chg_pct_hl = 0, chg_pct_sl = 0, chg_hl = 0, chg_sl = 0, chg = 0
                WHERE fp.period_name = ANY($1)
                  AND NOT EXISTS (
                      SELECT 1 FROM chargeability_blocks cb
                      WHERE cb.eid = fp.eid AND cb.period_name = fp.period_name
                  )
                """,
                periodos_db)

        log.info('chargeability_blocks y forecast_periods actualizados')

        restantes = await conn.fetch(
            """
            SELECT e.eid, e.offering AS emp_off, fu.offering AS fu_off
            FROM employees e
            LEFT JOIN (
                SELECT DISTINCT ON (eid) eid, offering FROM forecast_update
                ORDER BY eid, updated_at DESC NULLS LAST
            ) fu ON fu.eid = e.eid
            WHERE e.active AND COALESCE(e.offering, '') <> COALESCE(fu.offering, '')
            """)
        if restantes:
            log.warning(f'{len(restantes)} empleados quedaron desalineados entre '
                        'employees y forecast_update:')
            for r in restantes[:12]:
                log.warning(f"    {r['eid']:26} employees={r['emp_off']} "
                            f"forecast_update={r['fu_off']}")
        else:
            log.info('employees y forecast_update alineados')

        print()
        print('  Listo. Ahora corre: python create_daily_hours.py')
        print()
    finally:
        await conn.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--legacy', default=DEFAULT_LEGACY,
                    help='Forecast S&P 9-1 (.xlsx) — SO/PR/Tools')
    ap.add_argument('--digi', default=DEFAULT_DIGI,
                    help='01 - FORECAST CONTROL S&P (.xlsm) — S4/Ariba/Oracle')
    ap.add_argument('--min-period-col', type=int, default=LG_MIN_PERIOD_COL,
                    help=f'primera columna del bloque vigente del 9-1 '
                         f'(default {LG_MIN_PERIOD_COL}). Revisar el log '
                         '"[9-1] Periodos vigentes" para confirmar.')
    ap.add_argument('--check-period', default='Sep-P1',
                    help='periodo del reporte de totales (default Sep-P1)')
    ap.add_argument('--update-location', action='store_true',
                    help='tambien sincroniza location y country (afecta el filtro Pais)')
    ap.add_argument('--dry-run', action='store_true', help='no escribe en la DB')
    args = ap.parse_args()

    for etiqueta, ruta in (('9-1', args.legacy), ('CONTROL', args.digi)):
        if not os.path.exists(ruta):
            raise SystemExit(f'[{etiqueta}] No se encontro el archivo: {ruta}')

    if not args.dry_run and not all(
            [AZURE['host'], AZURE['user'], AZURE['password'], AZURE['database']]):
        raise SystemExit('Faltan variables de DB en .env')

    emp_lg, dat_lg, per_lg = parse_legacy(args.legacy, args.min_period_col)
    emp_dg, dat_dg, per_dg = parse_digi(args.digi)

    solapados = set(emp_lg) & set(emp_dg)
    if solapados:
        log.warning(f'{len(solapados)} EIDs aparecen como legacy y como Digi. '
                    f'Gana Digi (CONTROL): {sorted(solapados)}')

    # Digi manda para su gente; el 9-1 para el resto. Un EID sale de una sola
    # planilla, nunca de las dos, para que no queden horas mezcladas.
    empleados = {**emp_lg, **emp_dg}
    datos = defaultdict(dict)
    for eid, d in dat_lg.items():
        if eid not in emp_dg:
            datos[eid] = d
    for eid, d in dat_dg.items():
        datos[eid] = d

    log.info(f'Total combinado: {len(empleados)} empleados · '
             f'{Counter(v["fuente"] for v in empleados.values())}')

    periodos = {'9-1': per_lg, 'CONTROL': per_dg}
    check = args.check_period
    reporte(empleados, datos, check)

    asyncio.run(write_db(empleados, datos, periodos, args.dry_run, args.update_location))


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)