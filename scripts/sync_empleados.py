"""
alta_faltantes_9_1.py
---------------------
Da de alta en employees a la gente que esta en el Forecast S&P 9-1 y no en la
base, y normaliza el campo country.

Por que hace falta: el HC de la app no cierra contra el bloque de totales del
Excel (PR 12 vs 15, Tools 17 vs 18) porque esas personas no existen en
employees, asi que el importer las saltea. Y Mexico y Costa Rica aparecen en
cero porque no hay nadie con country MX ni CR.

El formato de country es el corto ('AR', 'MX', 'CR'): es lo que ya usan los 99
empleados existentes y lo que espera el filtro de Pais del front. Los dos altas
de ayer quedaron con 'Argentina' y se corrigen de paso.

Uso:
    python alta_faltantes_9_1.py --dry-run
    python alta_faltantes_9_1.py
"""

import argparse
import asyncio
import logging
import os
import sys
from collections import Counter

import asyncpg
from dotenv import load_dotenv
from openpyxl import load_workbook

# El .env vive en la raiz del repo, no en scripts/
load_dotenv(os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), '.env'))

logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s',
                    datefmt='%H:%M:%S')
log = logging.getLogger(__name__)

AZURE = dict(
    host=os.getenv('DB_HOST'),
    port=int(os.getenv('DB_PORT', 5432)),
    user=os.getenv('DB_USER'),
    password=os.getenv('DB_PASSWORD'),
    database=os.getenv('DB_NAME'),
    ssl='require',
)

DESKTOP = os.path.join(os.path.expanduser('~'), 'OneDrive - Accenture', 'Desktop')
DEFAULT_XLSX = os.path.join(DESKTOP, '2026 9 1 - Forecast S&P.xlsx')

SHEET = 'Forecast Update'
FIRST_DATA_ROW = 11
COL_LOCATION = 1
COL_EID = 2
COL_OFFERING = 3
COL_CL = 4

OFFERINGS = {'SO', 'PR', 'Tools'}
LOCATIONS = {'ARG', 'MX', 'CR'}

# La base usa el codigo corto. 'ARG' del Excel es 'AR' en employees.country.
LOC_A_COUNTRY = {'ARG': 'AR', 'MX': 'MX', 'CR': 'CR'}


def nombre_desde_eid(eid: str) -> str:
    """'ana.l.quintero' -> 'Ana L Quintero'. El 9-1 no trae el nombre completo;
    se arma desde el EID para que la lista de la app no muestre celdas vacias."""
    return ' '.join(p.capitalize() for p in eid.split('.') if p)


def parse_excel(path: str) -> dict:
    log.info(f'Leyendo {path}')
    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"No existe la hoja '{SHEET}'. Hojas: {wb.sheetnames}")
    ws = wb[SHEET]

    out = {}
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        eid = row[COL_EID - 1]
        if not isinstance(eid, str) or not eid.strip() or eid.startswith('='):
            continue
        loc = row[COL_LOCATION - 1]
        loc = loc.strip().upper() if isinstance(loc, str) else ''
        off = row[COL_OFFERING - 1]
        off = off.strip() if isinstance(off, str) else ''
        if loc not in LOCATIONS or off not in OFFERINGS:
            continue
        cl = row[COL_CL - 1]
        out[eid.strip().lower()] = {
            'location': loc,
            'country': LOC_A_COUNTRY[loc],
            'offering': off,
            'cl': int(cl) if isinstance(cl, (int, float)) else None,
        }
    wb.close()
    log.info(f'9-1: {len(out)} empleados legacy (ARG/MX/CR con SO/PR/Tools)')
    return out


async def run(xl: dict, dry_run: bool):
    conn = await asyncpg.connect(**AZURE)
    try:
        db = {r['eid'].strip().lower(): r for r in await conn.fetch(
            'SELECT eid, name, active, country, offering FROM employees')}

        nuevos, reactivar, fix_country = [], [], []
        for eid, d in sorted(xl.items()):
            row = db.get(eid)
            if row is None:
                nuevos.append((eid, d))
            elif not row['active']:
                reactivar.append((eid, d))

        # country largo -> corto, para que el filtro de Pais los agrupe bien
        for eid, row in db.items():
            c = row['country']
            if c and c not in ('AR', 'MX', 'CR'):
                corto = {'ARGENTINA': 'AR', 'MEXICO': 'MX', 'COSTA RICA': 'CR'}.get(
                    c.strip().upper())
                if corto:
                    fix_country.append((eid, c, corto))

        print()
        print('  === altas y correcciones ===')
        print(f'  A insertar          : {len(nuevos)}')
        print(f'  A reactivar         : {len(reactivar)}')
        print(f'  country a normalizar: {len(fix_country)}')
        print()

        if nuevos:
            print(f"  {'eid':26} {'loc':5} {'country':8} {'off':7} {'CL':>3}  nombre")
            print('  ' + '-' * 74)
            for eid, d in nuevos:
                print(f"  {eid:26} {d['location']:5} {d['country']:8} "
                      f"{d['offering']:7} {str(d['cl'] or '-'):>3}  {nombre_desde_eid(eid)}")
            print()
            print(f"  por pais: {dict(Counter(d['country'] for _e, d in nuevos))}")
            print()

        if reactivar:
            print('  --- a reactivar (existen pero inactivos) ---')
            for eid, d in reactivar:
                print(f"    {eid:26} {d['country']} {d['offering']}")
            print()

        if fix_country:
            print('  --- country a normalizar ---')
            for eid, a, b in fix_country:
                print(f'    {eid:26} {a}  ->  {b}')
            print()

        if dry_run:
            print('  DRY RUN: no se escribio nada en la DB.')
            print()
            return

        if not (nuevos or reactivar or fix_country):
            log.info('Nada que hacer.')
            return

        async with conn.transaction():
            if nuevos:
                await conn.executemany(
                    """
                    INSERT INTO employees
                        (eid, name, location, country, cl, offering,
                         active, charge, ringfenced, fte)
                    VALUES ($1, $2, $3, $4, $5, $6, TRUE, TRUE, FALSE, 1)
                    """,
                    [(eid, nombre_desde_eid(eid), d['location'], d['country'],
                      d['cl'], d['offering']) for eid, d in nuevos])

            if reactivar:
                await conn.executemany(
                    'UPDATE employees SET active = TRUE, offering = $2, '
                    'country = $3, location = $4 WHERE eid = $1',
                    [(eid, d['offering'], d['country'], d['location'])
                     for eid, d in reactivar])

            if fix_country:
                await conn.executemany(
                    'UPDATE employees SET country = $2 WHERE eid = $1',
                    [(eid, b) for eid, _a, b in fix_country])

            # forecast_update es la tabla por la que agrupan los totales. Sin
            # fila ahi el empleado cae en 'SIN OFFERING' y su SAH no entra al
            # denominador.
            todos = [(eid, d['offering']) for eid, d in nuevos + reactivar]
            if todos:
                existentes = {r['eid'] for r in await conn.fetch(
                    'SELECT DISTINCT eid FROM forecast_update WHERE eid = ANY($1)',
                    [e for e, _o in todos])}
                faltan = [(e, o) for e, o in todos if e not in existentes]
                if faltan:
                    await conn.executemany(
                        'INSERT INTO forecast_update (eid, offering, updated_at) '
                        'VALUES ($1, $2, NOW())', faltan)
                    log.info(f'{len(faltan)} filas nuevas en forecast_update')
                ya = [(e, o) for e, o in todos if e in existentes]
                if ya:
                    await conn.executemany(
                        'UPDATE forecast_update SET offering = $2, updated_at = NOW() '
                        'WHERE eid = $1', ya)

        log.info(f'{len(nuevos)} insertados · {len(reactivar)} reactivados · '
                 f'{len(fix_country)} country corregidos')

        r = await conn.fetch(
            'SELECT country, COUNT(*) n FROM employees WHERE active '
            'GROUP BY 1 ORDER BY 2 DESC')
        print()
        print('  empleados activos por pais:')
        for x in r:
            print(f"    {str(x['country']):8} {x['n']}")
        print()
        print('  Ahora corre:')
        print('    python import_forecast_combinado.py')
        print('    python create_daily_hours.py')
        print()
    finally:
        await conn.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--excel', default=DEFAULT_XLSX, help='Forecast S&P 9-1 (.xlsx)')
    ap.add_argument('--dry-run', action='store_true', help='no escribe en la DB')
    args = ap.parse_args()

    if not os.path.exists(args.excel):
        raise SystemExit(f'No se encontro el archivo: {args.excel}')
    if not args.dry_run and not all(
            [AZURE['host'], AZURE['user'], AZURE['password'], AZURE['database']]):
        raise SystemExit('Faltan variables de DB en .env')

    asyncio.run(run(parse_excel(args.excel), args.dry_run))


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)